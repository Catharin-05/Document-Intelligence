"""
Review PDF: one page per judged item, for checking the model by hand.

  py review.py build          -> <out_dir>/review.pdf   (one page per row of findings.xlsx)
  (fill in the PDF: "Your verdict" dropdown + "Your notes" box on every page, then save it)
  py review.py import         -> writes those fields into the human_verdict / human_note columns of findings.xlsx
  py judge.py export          -> results.corrected.json and findings.xlsx use your verdicts
review.py only imports judge.py; it does not change it.

Each page: the sheet with the item's regions in red and numbered (left) and a panel (right) with the item, the
model's verdict and what happens to it (REMOVED / kept), category, confidence, who decided it, the reason, the
region texts, and your fields. Nothing here calls a model.
"""
import json
import math
import re
from pathlib import Path

import pymupdf as fitz

GREEN, RED, GREY, AMBER, INK = (0.10, 0.50, 0.20), (0.78, 0.15, 0.15), (0.40, 0.40, 0.40), (0.80, 0.50, 0.0), (0.1, 0.1, 0.1)
PANEL = 560


def _latin(s):
    return str(s).encode("latin-1", "replace").decode("latin-1")


def _fieldname(prefix, stem):
    return f"{prefix}__{stem}".replace(".", "_")


def xlsx_file(cfg):
    """The spreadsheet to use: --xlsx path (review.xlsx in config), else <out_dir>/findings.xlsx."""
    return Path((cfg.get("review") or {}).get("xlsx") or Path(cfg["paths"]["out_dir"]) / "findings.xlsx")


def rc_source(cfg):
    return (cfg.get("review") or {}).get("source", "excel")


def _ints(x):
    return [int(t) for t in re.findall(r"\d+", str(x or ""))]


def from_excel(path, sheets, J, human):
    """Verdict rows of findings.xlsx -> {key: verdict dict}; also fills `human` from its human_verdict/human_note."""
    path = Path(path)
    if not path.exists():
        return {}
    from openpyxl import load_workbook
    try:
        ws = load_workbook(path, data_only=True)["Findings"]
    except PermissionError:
        raise SystemExit(f"{path.name} is open in Excel and cannot be read: close it and run again.")
    head = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(head) if h}
    by_name = {}
    for sh, pkg, idx, it in J.iter_items(sheets, ()):
        by_name[(sh["sheet_number"], pkg, (it.get("bid_item") or "").strip())] = J.item_key(sh, pkg, idx)
    out, missing = {}, 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        g = lambda name, d="": row[col[name]] if name in col and row[col[name]] is not None else d
        key = g("key") or by_name.get((g("sheet_number"), g("scope_package"), str(g("item_name")).strip()))
        if not key or not g("verdict"):
            missing += 1
            continue
        conf = g("confidence", 0.0)
        out[key] = {"key": key, "verdict": str(g("verdict")).lower(), "category": g("category", "other"),
                    "reason": str(g("reason")), "confidence": float(conf or 0),
                    "needs_human_review": str(g("needs_review")).lower() in ("yes", "true", "1"),
                    "supporting_regions": _ints(g("supporting_regions")), "unsupported_regions": _ints(g("unsupported_regions")),
                    "decided_by": g("decided_by", "image") or "image", "model": g("model")}
        hv, hn = str(g("human_verdict")).strip().lower(), str(g("human_note")).strip()
        if hv in ("right", "wrong") or hn:
            human[key] = {"verdict": hv if hv in ("right", "wrong") else "", "note": hn}
    if missing:
        print(f"review: {missing} spreadsheet rows skipped (no key / no verdict)")
    print(f"review: {len(out)} judged items read from {path.name}")
    return out


def build(cfg, J, only=None):
    """J = the judge module (helpers). only: None | 'wrong' | 'review'."""
    out = Path(cfg["paths"]["out_dir"])
    data, sheets = J.load(cfg["paths"]["results"])
    doc_src = fitz.open(cfg["paths"]["pdf"])
    human = {}
    verdicts = from_excel(xlsx_file(cfg), sheets, J, human) if rc_source(cfg) == "excel" else {}
    if not verdicts:  # no spreadsheet (or source: cache): use the saved verdicts
        for f in (out / "cache").glob("*.json"):
            j = json.loads(f.read_text(encoding="utf-8"))
            verdicts[j["key"]] = j

    thr = cfg["export"]["review_confidence_below"]
    rc = cfg.get("review", {})
    img_w, img_q = rc.get("image_width", 2400), rc.get("jpeg_quality", 72)

    pdf = fitz.open()
    toc, ctxs, n = [], {}, 0
    last_sheet = None
    for sh, pkg, idx, it in J.iter_items(sheets, cfg["legend"]["sheets"]):
        key = J.item_key(sh, pkg, idx)
        v = verdicts.get(key)
        if not v:
            continue
        flagged = v["needs_human_review"] or v["confidence"] < thr
        if (only == "wrong" and v["verdict"] != "wrong") or (only == "review" and not (flagged or v["verdict"] == "wrong")):
            continue
        n += 1
        ctx = ctxs.get(sh["sheet_number"]) or ctxs.setdefault(sh["sheet_number"], J.SheetCtx(doc_src, sh))
        regs = J.item_regions(it)
        img = J.jpeg(J.overview_image(ctx, regs, img_w), img_q)
        from PIL import Image
        import io
        iw, ih = Image.open(io.BytesIO(img)).size
        pg_h = max(ih / (img_w / 1512), 900)  # the sheet is drawn 1512 pt wide
        page = pdf.new_page(width=1512 + PANEL, height=pg_h)
        page.insert_image(fitz.Rect(0, 0, 1512, ih / (img_w / 1512)), stream=img)
        stem = J.file_stem(key)
        x0, x1 = 1512 + 18, 1512 + PANEL - 18
        w = x1 - x0
        y = 18

        def put(text, size=10.5, color=INK, bold=False, gap=6):
            nonlocal y
            text = _latin(text)
            cpl = max(10, int(w / (size * 0.50)))
            lines = sum(max(1, math.ceil(len(p) / cpl)) for p in text.split("\n"))
            h = lines * size * 1.3 + 4
            for _ in range(8):  # insert_textbox draws nothing when the box is too small: grow it until it fits
                if page.insert_textbox(fitz.Rect(x0, y, x1, y + h), text, fontsize=size,
                                       fontname="hebo" if bold else "helv", color=color) >= 0:
                    break
                h += size * 0.9
            y += h + gap

        page.draw_line((1512, 0), (1512, pg_h), color=(0.7, 0.7, 0.7), width=1.5)
        put(f"{sh['sheet_number']}  -  {sh.get('sheet_name', '')}   (pdf page {sh['page_number']}, item {idx + 1} of '{pkg}')", 9.5, GREY)
        put(f"SCOPE PACKAGE: {pkg}", 11, (0.0, 0.27, 0.55), True, 2)
        put(it.get("bid_item") or "", 14, INK, True, 2)
        if it.get("subheading"):
            put(it["subheading"], 10, GREY)
        wrong = v["verdict"] == "wrong"
        h_hv = human.get(key, {})
        put(("MODEL: WRONG  ->  REMOVED from results.json" if wrong else "MODEL: RIGHT  ->  kept, not modified"), 13,
            RED if wrong else GREEN, True, 2)
        put(f"category: {v['category']}    confidence: {v['confidence']:.2f}    decided by: {v.get('decided_by', 'image')}"
            f" ({v.get('model', '')})" + ("    NEEDS REVIEW" if flagged else ""), 9.5, AMBER if flagged else GREY)
        if v.get("escalated_from"):
            put(f"escalated from {v['escalated_from']['model']}: {v['escalated_from']['verdict']} "
                f"({v['escalated_from']['confidence']:.2f})", 9, GREY)
        put("REASON", 8.5, GREY, True, 1)
        put(v["reason"], 10.5)
        # regions
        groups = J.group_regions(regs)
        sup, uns = set(v.get("supporting_regions") or []), set(v.get("unsupported_regions") or [])
        put(f"REGIONS ({len(regs)} on the sheet, {len(groups)} distinct texts; numbers match the red labels)", 8.5, GREY, True, 1)
        shown = 0
        for t, nums, first, members in groups[:22]:
            tag = "+" if sup & set(nums) else ("-" if uns & set(nums) else " ")
            ns = ",".join(f"#{k}" for k in nums[:6]) + (f"+{len(nums) - 6}" if len(nums) > 6 else "")
            put(f"[{tag}] {ns}  x{len(nums)}  \"{t[:140]}\"", 9, INK, False, 1)
            shown += 1
        if len(groups) > shown:
            put(f"... and {len(groups) - shown} more distinct texts", 9, GREY)
        put("[+] model says supports   [-] model says does not support", 8, GREY, False, 8)
        # human fields
        fy = max(y + 4, pg_h - 250)
        page.insert_textbox(fitz.Rect(x0, fy, x1, fy + 14), "YOUR VERDICT (right / wrong)", fontsize=9, fontname="hebo", color=GREY)
        wv = fitz.Widget()
        wv.field_type = fitz.PDF_WIDGET_TYPE_COMBOBOX
        wv.field_name = _fieldname("v", stem)
        wv.choice_values = ["", "right", "wrong"]
        wv.field_value = h_hv.get("verdict", "") or ""
        wv.rect = fitz.Rect(x0, fy + 16, x0 + 170, fy + 38)
        wv.text_fontsize = 11
        wv.border_color, wv.fill_color = (0.4, 0.4, 0.4), (1, 1, 0.85)
        page.add_widget(wv)
        page.insert_textbox(fitz.Rect(x0, fy + 46, x1, fy + 60), "YOUR NOTES", fontsize=9, fontname="hebo", color=GREY)
        wn = fitz.Widget()
        wn.field_type = fitz.PDF_WIDGET_TYPE_TEXT
        wn.field_name = _fieldname("n", stem)
        wn.field_flags = fitz.PDF_TX_FIELD_IS_MULTILINE
        wn.field_value = h_hv.get("note", "") or ""
        wn.rect = fitz.Rect(x0, fy + 62, x1, pg_h - 18)
        wn.text_fontsize = 10
        wn.border_color, wn.fill_color = (0.4, 0.4, 0.4), (1, 1, 0.85)
        page.add_widget(wn)
        if sh["sheet_number"] != last_sheet:
            toc.append([1, f"{sh['sheet_number']}  {sh.get('sheet_name', '')}"[:100], n])
            last_sheet = sh["sheet_number"]
        toc.append([2, f"{'WRONG' if wrong else 'right'}{' *' if flagged else ''}  [{pkg}] {it.get('bid_item') or ''}"[:110], n])
    if not n:
        print("review: nothing to show (no judged items match)")
        return None
    pdf.set_toc(toc)
    path = out / ("review.pdf" if not only else f"review_{only}.pdf")
    try:
        pdf.save(path, deflate=True, garbage=3)
    except Exception as e:
        path = path.with_name(path.stem + "_new.pdf")
        pdf.save(path, deflate=True, garbage=3)
        print(f"warning: could not overwrite the old file ({e}); saved as {path.name}")
    print(f"review: {n} pages -> {path}\n  fill in 'Your verdict' and 'Your notes' on each page, save the PDF, then run: py review.py import")
    return path


def import_fields(cfg, J):
    """Read the form fields of review*.pdf back into the human_verdict / human_note columns of findings.xlsx."""
    from openpyxl import load_workbook
    out = Path(cfg["paths"]["out_dir"])
    key_of = {}
    for f in (out / "cache").glob("*.json"):
        j = json.loads(f.read_text(encoding="utf-8"))
        key_of[J.file_stem(j["key"]).replace(".", "_")] = j["key"]
    _, sheets = J.load(cfg["paths"]["results"])  # also map keys of items that have no cache file (spreadsheet-only runs)
    for sh, pkg, idx, it in J.iter_items(sheets, ()):
        k = J.item_key(sh, pkg, idx)
        key_of.setdefault(J.file_stem(k).replace(".", "_"), k)
    src = (cfg.get("review") or {}).get("file") or None
    pdfs = [Path(src)] if src else sorted(out.glob("review*.pdf"), key=lambda p: p.stat().st_mtime)[-1:]
    mine = {}
    for p in pdfs:
        for page in fitz.open(p):
            for w in page.widgets() or []:
                m = re.match(r"^([vn])__(.+)$", w.field_name or "")
                if not m or m.group(2) not in key_of:
                    continue
                rec = mine.setdefault(key_of[m.group(2)], {})
                val = (w.field_value or "").strip()
                if m.group(1) == "v":
                    rec["verdict"] = val.lower() if val.lower() in ("right", "wrong") else ""
                else:
                    rec["note"] = val
        print(f"import: read {p.name}")
    xp = xlsx_file(cfg)
    if not xp.exists():
        raise SystemExit(f"{xp} not found")
    wb = load_workbook(xp)
    ws = wb["Findings"]
    head = [c.value for c in ws[1]]
    for name in ("human_verdict", "human_note"):
        if name not in head:
            ws.cell(1, len(head) + 1).value = name
            head.append(name)
    ki, vi, ni = head.index("key") + 1, head.index("human_verdict") + 1, head.index("human_note") + 1
    n_v = n_n = 0
    for r in range(2, ws.max_row + 1):
        rec = mine.get(ws.cell(r, ki).value)
        if not rec:
            continue
        if rec.get("verdict"):
            ws.cell(r, vi).value = rec["verdict"]
            n_v += 1
        if rec.get("note"):
            ws.cell(r, ni).value = rec["note"]
            n_n += 1
    try:
        wb.save(xp)
    except PermissionError:
        raise SystemExit(f"{xp.name} is open in Excel: close it and run again.")
    print(f"import: wrote {n_v} verdicts and {n_n} notes into {xp.name}\n  now run: py judge.py export   (applies your verdicts to results.corrected.json)")


def main():
    """Standalone: py review.py build|import  [--config config.yaml] [--only wrong|review]. Uses judge.py as-is."""
    import argparse
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("command", choices=["build", "import"])
    ap.add_argument("--config", default="config.yaml")
    ap.add_argument("--set", action="append", default=[])
    ap.add_argument("--only", choices=["wrong", "review"])
    ap.add_argument("--xlsx", help="use this findings spreadsheet (e.g. output\\run5\\findings_204813.xlsx); the PDF is saved next to it")
    a = ap.parse_args()
    import judge as J
    cfg = J.load_config(a.config, a.set)
    if a.xlsx:
        x = Path(a.xlsx).resolve()
        cfg.setdefault("review", {})["xlsx"] = str(x)
        cfg["paths"]["out_dir"] = str(x.parent)
    if a.command == "build":
        build(cfg, J, a.only)
    else:
        import_fields(cfg, J)


if __name__ == "__main__":
    main()