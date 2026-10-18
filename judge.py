#!/usr/bin/env python3
"""
WyreAI assignment - judge every bid item with a vision-language model.

FLOW
  1. legend   extract the legend sheets' text (abbreviations + assembly types)  -> <out_dir>/legend.txt
  2. judge    for each bid item: ONE image (the sheet with the item's regions in red, numbered)
              + the item, the region texts/positions, text from the PDF text layer around them,
              and the legend text  ->  structured verdict (right | wrong, category, reason, confidence)
              cached per item in <out_dir>/cache/
  3. export   <out_dir>/results.corrected.json   (wrong items removed, `scopes` kept in sync, same shape)
              <out_dir>/findings.xlsx            (one row per judged item + Summary sheet)
  run         = legend (if missing) + judge + export          models = list models your key can call

EVERYTHING IS SET IN config.yaml (paths, model, prompts, filters). The prompt is one file: prompts/judge.md.
The API key comes from the environment:  GEMINI_API_KEY  (provider gemini)  or  DASHSCOPE_API_KEY  (qwen)  or  GROQ_API_KEY  (groq)

  python judge.py                                                           # same as: run
  python judge.py run     --config config.yaml
  python judge.py run     --config config.yaml --mock                       # no API; heuristic stub
  python judge.py judge   --config config.yaml --set filters.sheets=[A-421] --set filters.limit=3
  python judge.py legend  --config config.yaml                              # then read/fix legend.txt
  python judge.py export  --config config.yaml                              # after editing human_verdict in findings.xlsx
  python judge.py models

Cached verdicts are reused only if the prompt, legend and model are unchanged (so editing a prompt
re-judges automatically); --force re-judges everything.
pip install -r requirements.txt
"""
import argparse
from triage import LegendIndex, rule_verdict
import collections
import concurrent.futures as cf
import copy
import hashlib
import io
import json
import os
import re
import sys
import threading
import time
from pathlib import Path
from typing import List, Literal

import pymupdf as fitz
import yaml
from PIL import Image, ImageDraw, ImageFont
from pydantic import BaseModel

from legend import extract_legend_text

RED = (229, 57, 53)

DEFAULTS = {
    "paths": {"results": "data/results.json", "pdf": "data/drawings.pdf", "out_dir": "output/run1"},
    "legend": {"sheets": ["G-002.3", "G-501", "G-511"], "file": None, "max_chars": 60000},
    "triage": {"rules": True, "legend_auto": True, "trim_legend": True, "text_first": True,
               "short_text_len": 20, "text_min_confidence": 0.8,
               "url_frac": 0.8, "grid_min_regions": 4, "grid_frac": 0.9, "grid_edge_frac": 0.6},
    "model": {"provider": "gemini", "base_url": None, "api_key_env": None, "extra_body": None, "name": "gemini-3.8-flash", "fallback_names": [], "escalate_name": None, "escalate_below": 0.7,
              "temperature": 0, "workers": 4, "retries": 5, "min_seconds_between_calls": 0},
    "prompts": {"file": "prompts/judge.md"},
    "image": {"max_width": 3024, "jpeg_quality": 85},
    "context": {"max_regions_listed": 60, "nearby_text_regions": 6, "nearby_text": True},
    "filters": {"sheets": [], "packages": [], "limit": 0},
    "export": {"review_confidence_below": 0.7},
}


# ============================================================== config
def deep_merge(base, new):
    out = copy.deepcopy(base)
    for k, v in (new or {}).items():
        out[k] = deep_merge(out[k], v) if isinstance(v, dict) and isinstance(out.get(k), dict) else v
    return out


def load_config(path, overrides=()):
    p = Path(path).resolve()
    cfg = deep_merge(DEFAULTS, yaml.safe_load(p.read_text(encoding="utf-8")) if p.exists() else {})
    for o in overrides:  # --set a.b=value (value parsed as YAML: 5, true, [A-421], text ...)
        k, _, v = o.partition("=")
        node = cfg
        parts = k.strip().split(".")
        for part in parts[:-1]:
            node = node.setdefault(part, {})
        node[parts[-1]] = yaml.safe_load(v)
    base = p.parent
    cfg["_base"] = base
    for sect, key in (("paths", "results"), ("paths", "pdf"), ("paths", "out_dir"),
                      ("prompts", "file"), ("legend", "file")):
        v = cfg[sect].get(key)
        if v:
            cfg[sect][key] = str((base / v) if not Path(v).is_absolute() else Path(v))
    return cfg


# ============================================================== data helpers
def load(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    sheets = data["drawings_processing_results"] if isinstance(data, dict) else data
    return data, sheets


def qkey(r):
    return (tuple(round(v, 1) for v in r.get("quad_px", [])), r.get("text"))


def item_key(sh, pkg, idx):
    return f"{sh['sheet_number']}|{pkg}|{idx}"


def safe(s, n=60):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", s)[:n].strip("_")


def file_stem(key):
    sn, pkg, idx = key.split("|")
    return f"{safe(sn)}__{int(idx):03d}__{safe(pkg, 22)}__{hashlib.sha1(key.encode()).hexdigest()[:6]}"


def iter_items(sheets, skip=(), only_sheets=None, only_pkgs=None):
    for sh in sheets:
        if sh["sheet_number"] in skip:
            continue
        if only_sheets and sh["sheet_number"] not in only_sheets:
            continue
        for pkg, items in sh.get("bid_items", {}).items():
            if only_pkgs and pkg not in only_pkgs:
                continue
            for idx, it in enumerate(items):
                yield sh, pkg, idx, it


def item_regions(it):
    """Regions drawn on the sheet itself (the red ones). Deduplicated."""
    out, seen = [], set()
    for r in it.get("regions") or []:
        k = qkey(r)
        if len(k[0]) == 8 and k not in seen:
            seen.add(k)
            out.append(r)
    return out


# ============================================================== sheet context (render + text layer)
class SheetCtx:
    """Rendered page + PDF text layer, both in the 3024x2160 'as you see it' space."""

    def __init__(self, doc, sh):
        self.page = doc[sh["page_number"] - 1]
        self.W, self.H = sh.get("page_width", 3024), sh.get("page_height", 2160)
        self.sx, self.sy = self.W / self.page.rect.width, self.H / self.page.rect.height
        self._img = self._words = None

    @property
    def img(self):
        if self._img is None:
            pix = self.page.get_pixmap(matrix=fitz.Matrix(self.sx, self.sy), alpha=False)
            self._img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        return self._img

    @property
    def words(self):
        """(x0, y0, x1, y1, text) in displayed coords. get_text() works in the unrotated page
        space, so pages stored with /Rotate need page.rotation_matrix applied."""
        if self._words is None:
            ws = []
            for w in self.page.get_text("words"):
                r = fitz.Rect(w[:4]) * self.page.rotation_matrix
                r.normalize()
                ws.append((r.x0 * self.sx, r.y0 * self.sy, r.x1 * self.sx, r.y1 * self.sy, w[4]))
            self._words = ws
        return self._words

    def nearby_text(self, quad, rx=260, ry=110, max_words=24):
        xs, ys = quad[0::2], quad[1::2]
        x0, y0, x1, y1 = min(xs), min(ys), max(xs), max(ys)
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        near = []
        for wx0, wy0, wx1, wy1, t in self.words:
            if wx1 < x0 - rx or wx0 > x1 + rx or wy1 < y0 - ry or wy0 > y1 + ry:
                continue
            if wx0 >= x0 - 2 and wx1 <= x1 + 2 and wy0 >= y0 - 2 and wy1 <= y1 + 2:
                continue  # the region's own text
            if not t.strip("_-= ."):
                continue  # drawing rules / underlines
            near.append((abs((wx0 + wx1) / 2 - cx) + 2 * abs((wy0 + wy1) / 2 - cy), wy0, wx0, t))
        near.sort()
        near = sorted(near[:max_words], key=lambda n: (round(n[1] / 12), n[2]))
        return " ".join(n[3] for n in near)


def font(size, bold=False):
    for n in (["DejaVuSans-Bold.ttf", "arialbd.ttf"] if bold else ["DejaVuSans.ttf", "arial.ttf"]):
        try:
            return ImageFont.truetype(n, size)
        except OSError:
            pass
    return ImageFont.load_default()


def bbox(q, pad=0):
    xs, ys = q[0::2], q[1::2]
    return min(xs) - pad, min(ys) - pad, max(xs) + pad, max(ys) + pad


def overview_image(ctx, regs, max_w=3024):
    """The sheet with ONLY this item's regions drawn (red, numbered)."""
    base = ctx.img.convert("RGBA")
    layer = Image.new("RGBA", base.size, (0, 0, 0, 0))
    d = ImageDraw.Draw(layer)
    nf = font(20, True)
    for n, r in enumerate(regs, 1):
        q = r["quad_px"]
        pts = [(q[i], q[i + 1]) for i in range(0, 8, 2)]
        d.polygon(pts, fill=RED + (60,))
        d.line(pts + [pts[0]], fill=RED + (255,), width=3)
        x0, y0 = min(p[0] for p in pts), min(p[1] for p in pts)
        tw = d.textlength(str(n), font=nf)
        d.rectangle([x0, y0 - 24, x0 + tw + 8, y0], fill=RED + (255,))
        d.text((x0 + 4, y0 - 24), str(n), fill=(255, 255, 255, 255), font=nf)
    img = Image.alpha_composite(base, layer).convert("RGB")
    if img.width > max_w:
        img = img.resize((max_w, int(img.height * max_w / img.width)), Image.LANCZOS)
    return img


def group_regions(regs):
    """[(text, [region numbers], first_region, members)] in first-seen order; numbers match the image."""
    groups = collections.OrderedDict()
    for n, r in enumerate(regs, 1):
        t = " ".join((r.get("text") or "").split())
        groups.setdefault(t, []).append((n, r))
    return [(t, [n for n, _ in v], v[0][1], v) for t, v in groups.items()]


def pick_regions_for_context(groups, max_n=8):
    """Representative regions (largest text groups first) used for the 'text around it' lines."""
    picks = []
    for t, nums, first, members in sorted(groups, key=lambda g: -len(g[1])):
        if len(picks) >= max_n:
            break
        picks.append((members[0][0], members[0][1]))
        if len(members) >= 4 and len(picks) < max_n:
            picks.append((members[-1][0], members[-1][1]))
    return picks[:max_n]


def jpeg(img, q=85):
    b = io.BytesIO()
    img.save(b, "JPEG", quality=q)
    return b.getvalue()


# ============================================================== cheap signals (facts, not verdicts)
GRID_LIKE = re.compile(r"^[A-Z]{1,2}\d{0,2}$|^\d{1,2}$")
DIM_LIKE = re.compile(r"""(\d+\s*'\s*-?\s*\d*)|(\d+\s*(\"|MM|CM|SF)\b)|^\d+(\.\d+)?$""", re.I)
URL_LIKE = re.compile(r"(https?://|www\.|@\w+\.\w+|\.com\b)", re.I)


def signals(ctx, regs):
    n = len(regs)
    texts = [" ".join((r.get("text") or "").split()) for r in regs]
    edge = 0
    for r in regs:
        x0, y0, x1, y1 = bbox(r["quad_px"])
        cx, cy = (x0 + x1) / 2 / ctx.W, (y0 + y1) / 2 / ctx.H
        edge += (min(cx, 1 - cx) < 0.12) or (min(cy, 1 - cy) < 0.12)
    return {
        "n_regions": n,
        "n_distinct_texts": len(set(texts)),
        "frac_short_codes(<=3 chars)": round(sum(len(t) <= 3 for t in texts) / max(n, 1), 2),
        "frac_look_like_grid_labels": round(sum(bool(GRID_LIKE.match(t)) for t in texts) / max(n, 1), 2),
        "frac_dimension_like": round(sum(bool(DIM_LIKE.search(t)) and len(t) < 14 for t in texts) / max(n, 1), 2),
        "frac_url_or_email": round(sum(bool(URL_LIKE.search(t)) for t in texts) / max(n, 1), 2),
        "frac_within_12pct_of_sheet_edge": round(edge / max(n, 1), 2),
        "max_text_len": max((len(t) for t in texts), default=0),
    }


# ============================================================== prompts
class Verdict(BaseModel):
    verdict: Literal["right", "wrong"]
    category: Literal[
        "supported", "grid_bubble", "room_label", "dimension_or_annotation",
        "title_block_or_page_furniture", "abbreviation_misread", "wrong_trade_or_scope",
        "not_work", "other",
    ]
    reason: str
    supporting_regions: List[int]
    unsupported_regions: List[int]
    confidence: float
    needs_human_review: bool


def render(template, **fields):
    for k, v in fields.items():
        template = template.replace("{{" + k + "}}", str(v))
    return template


def build_prompt(template, cfg, sh, pkg, it, regs, groups, picks, ctx, sig):
    c = cfg["context"]
    extra = "".join(f"{f.upper()}: {it[f]}\n" for f in ("drawing_reference_detail", "notes") if it.get(f))
    lines = []
    for t, nums, first, members in groups[: c["max_regions_listed"]]:
        x0, y0, x1, y1 = bbox(first["quad_px"])
        pos = f"x={round((x0+x1)/2/ctx.W*100)}%, y={round((y0+y1)/2/ctx.H*100)}% of sheet"
        ns = ",".join(f"#{n}" for n in nums[:10]) + (f",+{len(nums)-10} more" if len(nums) > 10 else "")
        lines.append(f'  {ns}  x{len(nums)}  "{t[:300]}"   first at {pos}')
    if len(groups) > c["max_regions_listed"]:
        rest = [g[0][:40] for g in groups[c["max_regions_listed"]:c["max_regions_listed"] + 100]]
        lines.append(f"  ... {len(groups)-c['max_regions_listed']} more distinct texts: " + " | ".join(rest))
    near = ""
    if c["nearby_text"]:
        ls = []
        for n, r in picks[: c["nearby_text_regions"]]:
            nt = ctx.nearby_text(r["quad_px"])
            if nt:
                ls.append(f"  around #{n}: {nt}")
        if ls:
            near = "\nTEXT FOUND IN THE PDF AROUND SOME OF THESE REGIONS (PDF text layer, nearest first):\n" + "\n".join(ls) + "\n"
    return render(template,
                  sheet=f"{sh['sheet_number']} - {sh.get('sheet_name','')}   (discipline: {sh.get('discipline','')})",
                  scope_package=pkg, bid_item=it.get("bid_item", ""), subheading=it.get("subheading", ""),
                  extra_fields=extra, n_regions=len(regs), n_distinct=len(groups),
                  regions="\n".join(lines), nearby_text=near, signals=json.dumps(sig))


# ============================================================== model calls
_client = None


class FatalModelError(RuntimeError):
    """Errors retrying cannot fix: wrong model name, bad key, no access, bad request."""


class QuotaExhausted(RuntimeError):
    """The model's DAILY quota is used up (retrying today cannot help)."""


_throttle = {"lock": threading.Lock(), "last": 0.0}


def throttle(min_gap):
    if not min_gap:
        return
    with _throttle["lock"]:
        wait = _throttle["last"] + min_gap - time.time()
        if wait > 0:
            time.sleep(wait)
        _throttle["last"] = time.time()


def retry_delay_seconds(msg):
    m = re.search(r"retry(?:Delay)?['\"]?[: ]+(?:in )?['\"]?([\d.]+)s", msg)
    return float(m.group(1)) if m else None


FATAL_MARKERS = ("NOT_FOUND", "404", "PERMISSION_DENIED", "403", "UNAUTHENTICATED", "401",
                 "API key", "INVALID_ARGUMENT", "400", "no longer available")


def gemini_call(model, parts, temperature=0, retries=5, min_gap=0):
    global _client
    from google import genai
    from google.genai import types
    if _client is None:
        _client = genai.Client()  # GEMINI_API_KEY / GOOGLE_API_KEY
    contents = [types.Part.from_bytes(data=p, mime_type="image/jpeg") if isinstance(p, bytes) else p for p in parts]
    cfg = types.GenerateContentConfig(response_mime_type="application/json",
                                      response_schema=Verdict, temperature=temperature)
    err = None
    for a in range(retries):
        try:
            throttle(min_gap)
            resp = _client.models.generate_content(model=model, contents=contents, config=cfg)
            v = resp.parsed if getattr(resp, "parsed", None) else Verdict.model_validate_json(resp.text)
            return v.model_dump()
        except Exception as e:  # rate limit, 5xx, malformed json ...
            s = str(e)
            if "429" in s or "RESOURCE_EXHAUSTED" in s:
                delay = retry_delay_seconds(s)
                if "PerDay" in s or (delay is not None and delay > 120):
                    raise QuotaExhausted(s) from e  # a daily cap: waiting a few seconds will not help
                err = e
                time.sleep(min(90, delay if delay else 2 ** a * 2))  # per-minute limit: wait and retry
                continue
            if any(m in s for m in FATAL_MARKERS):
                raise FatalModelError(s) from e
            err = e
            time.sleep(min(60, 2 ** a * 2))
    raise RuntimeError(f"model call failed: {err}")


# ---- OpenAI-compatible provider (Qwen via Alibaba DashScope, or OpenRouter, vLLM, ...)
PROVIDER = {"provider": "gemini", "base_url": None, "api_key_env": None, "extra_body": None}
QWEN_URLS = {"intl": "https://dashscope-intl.aliyuncs.com/compatible-mode/v1",
             "cn": "https://dashscope.aliyuncs.com/compatible-mode/v1"}
_oai = None


def set_provider(mcfg):
    PROVIDER.update({k: mcfg.get(k) for k in PROVIDER})
    if PROVIDER["provider"] == "qwen":
        PROVIDER["base_url"] = QWEN_URLS.get(PROVIDER["base_url"] or "intl", PROVIDER["base_url"])
        PROVIDER["api_key_env"] = PROVIDER["api_key_env"] or "DASHSCOPE_API_KEY"
    elif PROVIDER["provider"] == "groq":
        PROVIDER["base_url"] = PROVIDER["base_url"] or "https://api.groq.com/openai/v1"
        PROVIDER["api_key_env"] = PROVIDER["api_key_env"] or "GROQ_API_KEY"
    elif PROVIDER["provider"] == "openai":
        PROVIDER["api_key_env"] = PROVIDER["api_key_env"] or "OPENAI_API_KEY"  # local server: not needed


def oai_client():
    global _oai
    if _oai is None:
        from openai import OpenAI
        key = os.environ.get(PROVIDER["api_key_env"] or "")
        if not key and PROVIDER["provider"] == "openai" and "localhost" in (PROVIDER["base_url"] or "") or \
                not key and "127.0.0.1" in (PROVIDER["base_url"] or ""):
            key = "local"  # local servers (Ollama, vLLM, LM Studio) ignore the key
        if not key:
            raise FatalModelError(f"set the environment variable {PROVIDER['api_key_env']}")
        _oai = OpenAI(api_key=key, base_url=PROVIDER["base_url"], timeout=300)
    return _oai


def parse_verdict_text(text):
    t = (text or "").strip()
    t = re.sub(r"^```(?:json)?|```$", "", t, flags=re.M).strip()
    t = re.sub(r"<think>.*?</think>", "", t, flags=re.S).strip()
    i, j = t.find("{"), t.rfind("}")
    return Verdict.model_validate_json(t[i:j + 1])


def openai_call(model, parts, temperature=0, retries=5, min_gap=0):
    import base64
    try:
        import openai  # noqa: F401
    except ImportError:
        raise FatalModelError("provider qwen/openai needs:  pip install openai")
    content = []
    for p in parts:
        if isinstance(p, bytes):
            content.append({"type": "image_url",
                            "image_url": {"url": "data:image/jpeg;base64," + base64.b64encode(p).decode()}})
        else:
            content.append({"type": "text", "text": p})
    content.append({"type": "text", "text": "Reply with ONE JSON object only (no markdown), matching this JSON schema:\n"
                    + json.dumps(Verdict.model_json_schema())})
    err = None
    for a in range(retries):
        try:
            throttle(min_gap)
            kw = {"extra_body": PROVIDER["extra_body"]} if PROVIDER["extra_body"] else {}
            resp = oai_client().chat.completions.create(
                model=model, messages=[{"role": "user", "content": content}], temperature=temperature,
                response_format={"type": "json_object"}, **kw)
            return parse_verdict_text(resp.choices[0].message.content).model_dump()
        except FatalModelError:
            raise
        except Exception as e:
            s = str(e)
            low = s.lower()
            if "429" in s or "rate" in low and "limit" in low or "quota" in low or "arrearage" in low or "insufficient" in low:
                if any(w in low for w in ("per day", "daily", "quota", "arrearage", "insufficient", "balance")):
                    raise QuotaExhausted(s) from e
                err = e
                time.sleep(min(90, 2 ** a * 3))
                continue
            if any(m in s for m in ("401", "403", "404", "400", "invalid_api_key", "model_not_found", "does not exist",
                                    "not support")):
                raise FatalModelError(s + "\n(Qwen note: the model must be a VISION model, e.g. a -vl model; text-only models reject images)") from e
            err = e
            time.sleep(min(60, 2 ** a * 2))
    raise RuntimeError(f"model call failed: {err}")


def call_model(model, parts, temperature=0, retries=5, min_gap=0):
    fn = gemini_call if PROVIDER["provider"] == "gemini" else openai_call
    return fn(model, parts, temperature, retries, min_gap)


class ModelPool:
    """Primary model first; when its daily quota is gone, move on to the fallback models."""

    def __init__(self, names):
        self.names, self.i, self.lock = [n for n in names if n], 0, threading.Lock()

    @property
    def current(self):
        return self.names[self.i] if self.i < len(self.names) else None

    def used_up(self, name):
        with self.lock:
            if self.current == name:
                nxt = self.names[self.i + 1] if self.i + 1 < len(self.names) else None
                print(f"\nDaily quota used up for {name}" + (f" -> switching to {nxt}" if nxt else " -> no more models to try"),
                      file=sys.stderr)
                self.i += 1


def mock_call(sig, regs):
    """Offline stub so the whole pipeline can be exercised without an API key. NOT a real judgement."""
    grid = sig["frac_look_like_grid_labels"] >= 0.9 and sig["frac_within_12pct_of_sheet_edge"] >= 0.6
    url = sig["frac_url_or_email"] >= 0.5
    dim = sig["frac_dimension_like"] >= 0.9
    if grid or url or dim:
        cat = "grid_bubble" if grid else ("title_block_or_page_furniture" if url else "dimension_or_annotation")
        return Verdict(verdict="wrong", category=cat, reason="[mock] heuristic: " + cat, supporting_regions=[],
                       unsupported_regions=list(range(1, len(regs) + 1)), confidence=0.6,
                       needs_human_review=True).model_dump()
    return Verdict(verdict="right", category="supported", reason="[mock] no heuristic flag raised",
                   supporting_regions=list(range(1, len(regs) + 1)), unsupported_regions=[], confidence=0.5,
                   needs_human_review=True).model_dump()


# ============================================================== stage 1: legend
def legend_path(cfg):
    return Path(cfg["legend"]["file"]) if cfg["legend"]["file"] else Path(cfg["paths"]["out_dir"]) / "legend.txt"


def cmd_legend(cfg, a=None):
    _, sheets = load(cfg["paths"]["results"])
    wanted = cfg["legend"]["sheets"]
    legend_sheets = [(s["sheet_number"], s.get("sheet_name", ""), s["page_number"]) for s in sheets if s["sheet_number"] in wanted]
    missing = set(wanted) - {s[0] for s in legend_sheets}
    if missing:
        print(f"warning: legend sheets not found in results.json: {sorted(missing)}", file=sys.stderr)
    text = extract_legend_text(cfg["paths"]["pdf"], sheets=legend_sheets)
    out = Path(cfg["paths"]["out_dir"]) / "legend.txt"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8")
    n_types = len(re.findall(r"^\[", text, re.M))
    print(f"legend: {len(text)} chars from {len(legend_sheets)} sheets, {n_types} assembly types -> {out}")
    print("        read it once and fix anything odd; to use a corrected copy set legend.file in config.yaml")
    return text


def get_legend(cfg):
    p = legend_path(cfg)
    if not p.exists():
        cmd_legend(cfg)
    text = p.read_text(encoding="utf-8")
    if len(text) > cfg["legend"]["max_chars"]:
        print(f"warning: legend is {len(text)} chars, cut to legend.max_chars={cfg['legend']['max_chars']}", file=sys.stderr)
        text = text[: cfg["legend"]["max_chars"]]
    return text


# ============================================================== stage 2: judge
def cmd_judge(cfg, a):
    pcfg, mcfg, icfg = cfg["paths"], cfg["model"], cfg["image"]
    set_provider(mcfg)
    _, sheets = load(pcfg["results"])
    doc = fitz.open(pcfg["pdf"])
    out = Path(pcfg["out_dir"])
    for d in ("cache", "evidence", "prompts"):
        (out / d).mkdir(parents=True, exist_ok=True)
    legend = get_legend(cfg)
    tpl = Path(cfg["prompts"]["file"]).read_text(encoding="utf-8")
    model = "mock-heuristic" if a.mock else mcfg["name"]
    tcfg = cfg["triage"]
    lidx = LegendIndex(legend)
    prompt_sha = hashlib.sha1((tpl + legend + model + str(mcfg["escalate_name"]) + json.dumps(tcfg, sort_keys=True)).encode()).hexdigest()[:10]
    IMG_NOTE = "The attached image is the sheet with this item's regions in red; the numbers match the list above. Give your verdict."
    TXT_NOTE = ("NO IMAGE is attached in this first pass, so you cannot see where the text sits or what it is attached to. "
                "Decide from the text, the legend lines and the facts above. If the answer depends on position or on what the "
                "text is attached to (a short tag, a code that is not in the legend, text that might be a floating label), "
                "set needs_human_review=true and confidence at or below 0.5 so that an image pass looks at it. "
                "Give your verdict.")

    f = cfg["filters"]
    items = list(iter_items(sheets, cfg["legend"]["sheets"], f["sheets"] or None, f["packages"] or None))
    if f["limit"]:
        items = items[: f["limit"]]
    todo = len(items)
    ex = cf.ThreadPoolExecutor(max_workers=mcfg["workers"])
    pool = ModelPool([mcfg["name"]] + list(mcfg["fallback_names"] or []))
    stop = threading.Event()
    fut_map, failed, done, cached = {}, [], 0, 0
    stats = collections.Counter()
    ctxs = {}

    def finish(block=False):
        nonlocal done
        ready, _ = cf.wait(list(fut_map), return_when=cf.FIRST_COMPLETED if block else cf.ALL_COMPLETED,
                           timeout=None if block else 0)
        for fu in ready:
            key, stem, meta = fut_map.pop(fu)
            try:
                res = fu.result()
            except QuotaExhausted:
                stop.set()  # every model is out of daily quota: stop cleanly, keep what is done
                continue
            except FatalModelError as e:
                ex.shutdown(wait=False, cancel_futures=True)
                sys.exit(f"\nSTOPPED: the model call cannot succeed, so no point continuing.\n{e}\n"
                         f"Run `python judge.py models` to see which models your key can use, then set model.name in config.yaml.")
            except Exception as e:
                failed.append((key, str(e)))
                print(f"FAILED {key}: {e}", file=sys.stderr)
                continue
            (out / "cache" / f"{stem}.json").write_text(json.dumps({**meta, **res}, indent=1), encoding="utf-8")
            done += 1
            stats[res.get("decided_by", "image")] += 1
            if done % 10 == 0:
                print(f"  judged {done}/{todo}")

    for sh, pkg, idx, it in items:
        key = item_key(sh, pkg, idx)
        stem = file_stem(key)
        cpath = out / "cache" / f"{stem}.json"
        if cpath.exists() and not a.force:
            try:
                if json.loads(cpath.read_text(encoding="utf-8")).get("prompt_sha") == prompt_sha:
                    cached += 1
                    continue
            except Exception:
                pass
        regs = item_regions(it)
        ctx = ctxs.get(sh["sheet_number"]) or ctxs.setdefault(sh["sheet_number"], SheetCtx(doc, sh))
        groups = group_regions(regs)
        picks = pick_regions_for_context(groups, max(cfg["context"]["nearby_text_regions"], 1))
        sig = signals(ctx, regs)
        text = build_prompt(tpl, cfg, sh, pkg, it, regs, groups, picks, ctx, sig)
        texts = [g[0] for g in groups]
        item_text = " ".join([it.get("bid_item") or "", it.get("subheading") or "", pkg])
        legend_part = lidx.relevant(" ".join(texts), item_text) if tcfg["trim_legend"] else legend
        text = text.replace("{{legend}}", legend_part)
        (out / "prompts" / f"{stem}.txt").write_text(text.replace("{{image_note}}", IMG_NOTE), encoding="utf-8")
        meta = {"key": key, "sheet_number": sh["sheet_number"], "package": pkg, "index": idx,
                "bid_item": it.get("bid_item", ""), "model": model, "prompt_sha": prompt_sha,
                "signals": sig, "evidence": f"evidence/{stem}.jpg"}

        # ---- stage 0: no model needed
        pre = None
        if a.mock:
            pre = {**mock_call(sig, regs), "decided_by": "mock"}
        else:
            if tcfg["rules"]:
                pre = rule_verdict(sig, tcfg, texts)
            if pre is None and tcfg["legend_auto"]:
                pre = lidx.auto_right(texts, item_text)
        if pre is not None:
            if pre["verdict"] == "wrong":
                pre["unsupported_regions"] = list(range(1, len(regs) + 1))
            else:
                pre["supporting_regions"] = list(range(1, len(regs) + 1))
            ov = jpeg(overview_image(ctx, regs, icfg["max_width"]), icfg["jpeg_quality"])
            (out / "evidence" / f"{stem}.jpg").write_bytes(ov)
            cpath.write_text(json.dumps({**meta, "model": pre["decided_by"], **pre}, indent=1), encoding="utf-8")
            done += 1
            stats[pre["decided_by"]] += 1
            continue

        short = sig["max_text_len"] <= tcfg["short_text_len"]

        def work(text=text, ctx=ctx, regs=regs, icfg=icfg, stem=stem, short=short):
            def ask(m, parts):
                return call_model(m, parts, mcfg["temperature"], mcfg["retries"], mcfg["min_seconds_between_calls"])

            def with_pool(parts):
                while True:
                    m = pool.current
                    if m is None:
                        raise QuotaExhausted("all models are out of daily quota")
                    try:
                        r = ask(m, parts)
                        r["model"] = m
                        return r
                    except QuotaExhausted:
                        pool.used_up(m)

            if stop.is_set():
                raise QuotaExhausted("stopped")
            first = None
            if tcfg["text_first"] and not short:  # short tags (P-1, EP-1 ...) always need the picture
                first = with_pool([text.replace("{{image_note}}", TXT_NOTE)])
                if not first["needs_human_review"] and first["confidence"] >= tcfg["text_min_confidence"]:
                    first["decided_by"] = "text"
                    return first
            ov = jpeg(overview_image(ctx, regs, icfg["max_width"]), icfg["jpeg_quality"])
            (out / "evidence" / f"{stem}.jpg").write_bytes(ov)
            r = with_pool([ov, text.replace("{{image_note}}", IMG_NOTE)])
            r["decided_by"] = "image"
            if first:
                r["text_pass"] = {"verdict": first["verdict"], "confidence": first["confidence"]}
            if mcfg["escalate_name"] and (r["confidence"] < mcfg["escalate_below"] or r["needs_human_review"]):
                try:
                    r2 = ask(mcfg["escalate_name"], [ov, text.replace("{{image_note}}", IMG_NOTE)])
                    return {**r2, "model": mcfg["escalate_name"], "decided_by": "image",
                            "escalated_from": {"model": r["model"], "verdict": r["verdict"], "confidence": r["confidence"]}}
                except QuotaExhausted:
                    pass  # keep the first verdict, it is already flagged for review
            return r

        fut_map[ex.submit(work)] = (key, stem, meta)
        while len(fut_map) >= mcfg["workers"] * 2:
            finish(block=True)
        if stop.is_set():
            break
    while fut_map:
        finish(block=True)
    complete = done + cached == todo
    print(f"judge: {done} new verdicts, {cached} reused from cache, {len(failed)} failed (of {todo} items)")
    if stats:
        print("  decided by: " + ", ".join(f"{k} {v}" for k, v in stats.items())
              + "   (rule/legend = no model call; text = text-only call; image = call with the sheet image)")
    if stop.is_set() or not complete:
        print(f"\nINCOMPLETE: {done + cached} of {todo} items have a verdict. Nothing is lost: verdicts are cached.\n"
              f"  - daily quota gone? Run the same command later (tomorrow), turn on billing, or add models to model.fallback_names.\n"
              f"  - the corrected JSON and spreadsheet below only reflect the items judged so far; do not submit them yet.",
              file=sys.stderr)
    if failed:
        (out / "failed.json").write_text(json.dumps(failed, indent=1))
    return complete


# ============================================================== stage 3: export
def load_overrides(xlsx):
    from openpyxl import load_workbook
    ov = {}
    p = Path(xlsx)
    if not p.exists():
        return ov
    try:
        ws = load_workbook(p, data_only=True)["Findings"]
        head = [c.value for c in ws[1]]
        ki, vi, ni = head.index("key"), head.index("human_verdict"), head.index("human_note")
    except Exception:  # old/other layout, unreadable or no Findings sheet: there are no overrides to apply
        return ov
    for row in ws.iter_rows(min_row=2, values_only=True):
        v = (row[vi] or "").strip().lower() if row[vi] else ""
        if v in ("right", "wrong"):
            ov[row[ki]] = (v, row[ni] or "")
    return ov


def cmd_export(cfg, a=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    data, sheets = load(cfg["paths"]["results"])
    legend_sheets = cfg["legend"]["sheets"]
    out = Path(cfg["paths"]["out_dir"])
    thr = cfg["export"]["review_confidence_below"]
    xlsx_path = out / "findings.xlsx"
    overrides = load_overrides(xlsx_path)
    verdicts = {}
    for fpath in (out / "cache").glob("*.json"):
        j = json.loads(fpath.read_text(encoding="utf-8"))
        verdicts[j["key"]] = j

    rows, decision = [], {}
    for sh, pkg, idx, it in iter_items(sheets, legend_sheets):
        key = item_key(sh, pkg, idx)
        v = verdicts.get(key)
        if not v:
            continue
        regs = item_regions(it)
        texts = [" ".join((r.get("text") or "").split()) for r in regs]
        hv, hn = overrides.get(key, ("", ""))
        final = hv or v["verdict"]
        decision[key] = final
        rows.append({
            "key": key, "sheet_number": sh["sheet_number"], "page_number": sh["page_number"],
            "sheet_name": sh.get("sheet_name", ""), "scope_package": pkg, "item_name": it.get("bid_item", ""),
            "subheading": it.get("subheading", ""), "verdict": v["verdict"], "category": v["category"],
            "reason": v["reason"], "confidence": round(v["confidence"], 2),
            "needs_review": "yes" if (v["needs_human_review"] or v["confidence"] < thr) else "",
            "n_regions": len(regs), "n_distinct_texts": len(set(texts)),
            "sample_texts": " | ".join(list(dict.fromkeys(texts))[:8])[:300],
            "supporting_regions": ",".join(map(str, v.get("supporting_regions", []))),
            "unsupported_regions": ",".join(map(str, v.get("unsupported_regions", []))),
            "decided_by": v.get("decided_by", "image"), "model": v.get("model", ""), "escalated": "yes" if v.get("escalated_from") else "",
            "human_verdict": hv, "human_note": hn, "final_verdict": final,
            "action": "REMOVED" if final == "wrong" else "kept", "evidence_image": v.get("evidence", ""),
        })

    # ---- corrected results.json: same structure, wrong items removed, scopes kept in sync
    removed_q, kept_q = collections.defaultdict(set), collections.defaultdict(set)
    n_removed = 0
    for sh in sheets:
        if sh["sheet_number"] in legend_sheets:
            continue
        new_items = {}
        for pkg, items in sh.get("bid_items", {}).items():
            keep = []
            for idx, it in enumerate(items):
                key = item_key(sh, pkg, idx)
                allq = {qkey(r) for r in item_regions(it)} | {qkey(r) for s in it.get("specifications") or [] for r in s.get("regions") or []}
                if decision.get(key) == "wrong":
                    removed_q[(sh["sheet_number"], pkg)] |= allq
                    n_removed += 1
                else:
                    kept_q[(sh["sheet_number"], pkg)] |= allq
                    keep.append(it)
            if keep:
                new_items[pkg] = keep
        sh["bid_items"] = new_items
        new_scopes = []
        for sc in sh.get("scopes", []):
            k = (sh["sheet_number"], sc.get("scope"))
            before = sc.get("regions") or []
            sc["regions"] = [r for r in before if not (qkey(r) in removed_q[k] and qkey(r) not in kept_q[k])]
            if before and not sc["regions"]:
                continue
            new_scopes.append(sc)
        sh["scopes"] = new_scopes
    (out / "results.corrected.json").write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")

    # ---- spreadsheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Findings"
    cols = list(rows[0].keys()) if rows else ["key"]
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    widths = {"key": 26, "sheet_name": 26, "scope_package": 22, "item_name": 46, "subheading": 18, "reason": 70,
              "sample_texts": 40, "evidence_image": 36, "human_note": 30}
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)
        ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, i).fill = PatternFill("solid", fgColor="2F4F6F")
    red, grn, yel = (PatternFill("solid", fgColor=c) for c in ("F8D7DA", "DFF0D8", "FFF3CD"))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[cols.index("verdict")].fill = red if row[cols.index("verdict")].value == "wrong" else grn
        row[cols.index("final_verdict")].fill = red if row[cols.index("final_verdict")].value == "wrong" else grn
        if row[cols.index("needs_review")].value:
            row[cols.index("needs_review")].fill = yel
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet("Summary")
    s.append(["sheet", "judged", "right", "wrong", "needs_review"])
    by = collections.defaultdict(lambda: [0, 0, 0, 0])
    for r in rows:
        b = by[r["sheet_number"]]
        b[0] += 1
        b[1] += r["final_verdict"] == "right"
        b[2] += r["final_verdict"] == "wrong"
        b[3] += bool(r["needs_review"])
    for k, v in by.items():
        s.append([k] + v)
    s.append(["TOTAL"] + [sum(v[i] for v in by.values()) for i in range(4)])
    s.append([])
    s.append(["category (model verdict = wrong)", "count"])
    for k, v in collections.Counter(r["category"] for r in rows if r["verdict"] == "wrong").most_common():
        s.append([k, v])
    s.append([])
    s.append(["human overrides applied", len(overrides)])
    s.append(["models used", ", ".join(sorted({r["model"] for r in rows}))])
    s.column_dimensions["A"].width = 34
    try:
        wb.save(xlsx_path)
    except PermissionError:  # the file is open in Excel (Windows locks it)
        alt = xlsx_path.with_name(xlsx_path.stem + time.strftime("_%H%M%S") + xlsx_path.suffix)
        wb.save(alt)
        print(f"warning: {xlsx_path.name} is open in Excel, so it could not be overwritten. Saved as {alt.name}. "
              f"Close Excel to write the normal file.", file=sys.stderr)
        xlsx_path = alt
    print(f"export: {len(rows)} judged, {n_removed} removed -> {out/'results.corrected.json'} and {xlsx_path}")


def cmd_models(cfg=None, a=None):
    set_provider(cfg["model"])
    if PROVIDER["provider"] != "gemini":
        for m in oai_client().models.list():
            print(m.id)
        return
    from google import genai
    for m in genai.Client().models.list():
        acts = getattr(m, "supported_actions", None) or []
        if not acts or "generateContent" in acts:
            print(m.name)


# ============================================================== cli
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("command", nargs="?", default="run", choices=["legend", "judge", "export", "run", "models"],
                    help="default: run (legend + judge + export)")
    ap.add_argument("--config", default="config.yaml")
    ap.add_argument("--set", action="append", default=[], metavar="KEY=VALUE",
                    help="override a config value, e.g. --set model.name=gemini-3.8-flash --set filters.sheets=[A-421]")
    ap.add_argument("--mock", action="store_true", help="no API calls; heuristic stub for testing the plumbing")
    ap.add_argument("--force", action="store_true", help="re-judge items even if a valid cached verdict exists")
    a = ap.parse_args()
    cfg = load_config(a.config, a.set)
    if a.command == "models":
        return cmd_models()
    Path(cfg["paths"]["out_dir"]).mkdir(parents=True, exist_ok=True)
    if a.command == "legend":
        cmd_legend(cfg)
    elif a.command == "judge":
        cmd_judge(cfg, a)
    elif a.command == "export":
        cmd_export(cfg)
    else:
        get_legend(cfg)
        cmd_judge(cfg, a)
        cmd_export(cfg)


if __name__ == "__main__":
    main()